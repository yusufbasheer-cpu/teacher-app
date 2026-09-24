"""Fill editable PowerPoint slides while retaining their source design.

The source deck is never modified. Unsupported slides fail explicitly instead of
returning a deck with missing text or silently replacing its design.
"""

from __future__ import annotations

import copy
import io
import math
import re
from typing import Any
from zipfile import ZipFile, BadZipFile

from collections import Counter

from lxml import etree
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE_TYPE
from pptx.enum.text import MSO_ANCHOR, MSO_AUTO_SIZE
from pptx.oxml.ns import qn
from pptx.util import Pt


class TemplateIncompatible(Exception):
    def __init__(self, code: str, message: str, slide: int | None = None):
        super().__init__(message)
        self.code = code
        self.slide = slide


MAX_TEMPLATE_BYTES = 15 * 1024 * 1024
MAX_ZIP_ENTRIES = 2500
MAX_UNCOMPRESSED_BYTES = 100 * 1024 * 1024
MAX_PHYSICAL_SLIDES = 240
# Uploaded-template exports must never discard a generated section simply because the source
# design has a small content area. Sections are split into continuation slides instead.
MAX_PARTS_PER_SECTION = MAX_PHYSICAL_SLIDES
# Design-only templates (background art, no text areas) get text areas added
# for them. Their font size is chosen per section, so a longer section may
# use a smaller size and, at the limit, a few more slides.
MAX_PARTS_DESIGN_ONLY = MAX_PHYSICAL_SLIDES
DESIGN_BODY_SIZES = (22, 20, 18, 16, 14)
# Share of the content width kept for text when a lesson picture sits beside it.
IMAGE_TEXT_FACTOR = 0.58
REL_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"


def validate_source(data: bytes) -> None:
    if not data or len(data) > MAX_TEMPLATE_BYTES:
        raise TemplateIncompatible("INVALID_TEMPLATE", "The PowerPoint is empty or larger than 15 MB.")
    try:
        with ZipFile(io.BytesIO(data)) as archive:
            names = archive.namelist()
            if len(names) > MAX_ZIP_ENTRIES or sum(i.file_size for i in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise TemplateIncompatible("INVALID_TEMPLATE", "This PowerPoint exceeds the template processing limit.")
            if "ppt/presentation.xml" not in names or "[Content_Types].xml" not in names:
                raise TemplateIncompatible("INVALID_TEMPLATE", "This file is not a valid editable .pptx.")
            for name in names:
                if name.startswith("/") or ".." in name.replace("\\", "/").split("/"):
                    raise TemplateIncompatible("INVALID_TEMPLATE", "This PowerPoint contains an unsafe file path.")
            if any("vbaProject" in name or "embeddings/" in name for name in names):
                raise TemplateIncompatible("UNSUPPORTED_CONTENT", "This PowerPoint contains macros or embedded files that cannot be safely preserved.")
    except BadZipFile as exc:
        raise TemplateIncompatible("INVALID_TEMPLATE", "This file is not a valid .pptx.") from exc


def _text_shapes(slide):
    return [s for s in slide.shapes if s.has_text_frame and s.width > 0 and s.height > 0]


def _role(shape) -> str | None:
    if not shape.is_placeholder:
        return None
    kind = str(shape.placeholder_format.type).upper()
    if "TITLE" in kind or "CENTER_TITLE" in kind:
        return "title"
    if "BODY" in kind or "SUBTITLE" in kind or "OBJECT" in kind:
        return "body"
    return None


def _editable_slots(slide, lenient: bool = False):
    """Find the title and body areas of a source slide.

    `lenient` is used for design-only slides that Layah added text areas to:
    decorative groups and small design pictures are then allowed, because
    nothing on such a slide is replaced except the text areas we added.
    """
    unsupported = {
        MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.DIAGRAM,
        MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.GROUP,
        MSO_SHAPE_TYPE.IGX_GRAPHIC, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT,
        MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.WEB_VIDEO,
    }
    if lenient:
        unsupported.discard(MSO_SHAPE_TYPE.GROUP)
    if any(shape.shape_type in unsupported for shape in slide.shapes):
        raise TemplateIncompatible("UNSUPPORTED_CONTENT", "This slide contains a chart, group, table, video or embedded object that cannot be safely replaced.")
    slide_width = slide.part.package.presentation_part.presentation.slide_width
    slide_height = slide.part.package.presentation_part.presentation.slide_height
    for shape in slide.shapes:
        if lenient or shape.shape_type != MSO_SHAPE_TYPE.PICTURE:
            continue
        area_ratio = shape.width * shape.height / (slide_width * slide_height)
        if area_ratio < 0.12 and 0.15 * slide_height < shape.top < 0.8 * slide_height:
            raise TemplateIncompatible("AMBIGUOUS_MAPPING", "This slide has a small lesson picture that cannot be distinguished from its design.")
    shapes = _text_shapes(slide)
    title = next((s for s in shapes if _role(s) == "title"), None)
    body = next((s for s in shapes if _role(s) == "body" and s is not title), None)
    if title is None:
        title_candidates = [s for s in shapes if s.top < slide.part.package.presentation_part.presentation.slide_height * 0.48]
        title = min(title_candidates, key=lambda s: (s.top, -s.width), default=None)
    if body is None:
        candidates = [s for s in shapes if s is not title and s.top < slide.part.package.presentation_part.presentation.slide_height * 0.9]
        body = max(candidates, key=lambda s: s.width * s.height, default=None)
    if title is None or body is None or title is body:
        raise TemplateIncompatible("NO_EDITABLE_CONTENT", "A source slide needs editable title and content text areas.")
    if title.width < 400000 or body.width < 500000 or body.height < 200000:
        raise TemplateIncompatible("NO_EDITABLE_CONTENT", "The editable title or content area is too small.")
    # Other substantial text would leak an old lesson into the new deck. A
    # short footer/label near the bottom is treated as retained branding.
    for shape in shapes:
        if shape is title or shape is body or not shape.text.strip():
            continue
        if shape.top > slide.part.package.presentation_part.presentation.slide_height * 0.85 and len(shape.text.strip()) < 60:
            continue
        raise TemplateIncompatible("AMBIGUOUS_MAPPING", "This slide has extra editable text. Use a simpler template with title and content areas.")
    return title, body


_HARD_UNSUPPORTED = {
    MSO_SHAPE_TYPE.CHART, MSO_SHAPE_TYPE.DIAGRAM,
    MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.IGX_GRAPHIC,
    MSO_SHAPE_TYPE.LINKED_OLE_OBJECT, MSO_SHAPE_TYPE.MEDIA,
    MSO_SHAPE_TYPE.TABLE, MSO_SHAPE_TYPE.WEB_VIDEO,
}


def _shape_has_text(shape) -> bool:
    if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
        return any(_shape_has_text(child) for child in shape.shapes)
    return shape.has_text_frame and bool(shape.text_frame.text.strip())


def _is_design_only(slide) -> bool:
    """True for a slide that is purely artwork: no text anywhere, nothing we cannot keep."""
    if any(_shape_has_text(shape) for shape in slide.shapes):
        return False
    return not any(shape.shape_type in _HARD_UNSUPPORTED for shape in slide.shapes)


def _background_blobs(slide):
    """Full-bleed pictures behind the slide, innermost layer first: (bytes, left, top, width, height)."""
    prs = slide.part.package.presentation_part.presentation
    slide_area = prs.slide_width * prs.slide_height
    for holder in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        found = []
        for shape in holder.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE and shape.width * shape.height >= 0.85 * slide_area:
                found.append((shape.image.blob, shape.left, shape.top, shape.width, shape.height))
        for blip in holder._element.iter(qn("a:blip")):
            rid = blip.get(qn("r:embed"))
            if rid and rid in holder.part.rels and not holder.part.rels[rid].is_external:
                # A background fill blip stretches over the whole slide.
                in_bg = any(parent.tag == qn("p:bg") for parent in blip.iterancestors())
                if in_bg:
                    found.append((holder.part.related_part(rid).blob, 0, 0, prs.slide_width, prs.slide_height))
        if found:
            return found
    return []


def _analyse_artwork(blob: bytes) -> dict | None:
    """Locate the clear area of a background picture.

    Works on a small copy: finds the widest and tallest band around the centre
    that has no strong artwork in it, and the artwork's dominant colour. Faint
    watermarks are ignored on purpose because text may sit on top of them.
    """
    try:
        from PIL import Image
        image = Image.open(io.BytesIO(blob)).convert("RGBA")
    except Exception:
        return None
    image.thumbnail((240, 240))
    canvas = Image.new("RGBA", image.size, (255, 255, 255, 255))
    canvas.alpha_composite(image)
    rgb = canvas.convert("RGB")
    width, height = rgb.size
    if width < 8 or height < 8:
        return None
    pixels = rgb.load()
    quantised = Counter((r // 16, g // 16, b // 16) for y in range(height) for x in range(width) for r, g, b in [pixels[x, y]])
    bin_ = quantised.most_common(1)[0][0]
    bg = tuple(v * 16 + 8 for v in bin_)

    def is_ink(x, y):
        r, g, b = pixels[x, y]
        return max(abs(r - bg[0]), abs(g - bg[1]), abs(b - bg[2])) > 90

    raw_ink = [[is_ink(x, y) for x in range(width)] for y in range(height)]
    # Only artwork that reaches the slide edge frames the content area. Ink that
    # floats inside (a centred crest or watermark) is left for text to sit over.
    ink = [[False] * width for _ in range(height)]
    stack = [(x, y) for y in range(height) for x in range(width)
             if raw_ink[y][x] and (x in (0, width - 1) or y in (0, height - 1))]
    for x, y in stack:
        ink[y][x] = True
    while stack:
        x, y = stack.pop()
        for dx in (-1, 0, 1):
            for dy in (-1, 0, 1):
                nx, ny = x + dx, y + dy
                if 0 <= nx < width and 0 <= ny < height and raw_ink[ny][nx] and not ink[ny][nx]:
                    ink[ny][nx] = True
                    stack.append((nx, ny))
    col = [sum(1 for y in range(height) if ink[y][x]) / height for x in range(width)]
    clean_col = [c < 0.02 for c in col]
    centre = width // 2
    if not clean_col[centre]:
        return None
    left = centre
    while left > 0 and clean_col[left - 1]:
        left -= 1
    right = centre
    while right < width - 1 and clean_col[right + 1]:
        right += 1
    row = [sum(1 for x in range(left, right + 1) if ink[y][x]) / max(1, right - left + 1) for y in range(height)]
    clean_row = [r < 0.02 for r in row]
    middle = height // 2
    if not clean_row[middle]:
        return None
    top = middle
    while top > 0 and clean_row[top - 1]:
        top -= 1
    bottom = middle
    while bottom < height - 1 and clean_row[bottom + 1]:
        bottom += 1

    inked = [pixels[x, y] for y in range(height) for x in range(width) if ink[y][x]]
    accent = None
    if inked:
        accent_bin = Counter((r // 32, g // 32, b // 32) for r, g, b in inked).most_common(1)[0][0]
        members = [c for c in inked if (c[0] // 32, c[1] // 32, c[2] // 32) == accent_bin]
        accent = tuple(sum(c[i] for c in members) // len(members) for i in range(3))
    area = [pixels[x, y] for y in range(top, bottom + 1) for x in range(left, right + 1)]
    luminance = sum(0.2126 * r + 0.7152 * g + 0.0722 * b for r, g, b in area) / max(1, len(area))
    return {
        "box": (left / width, top / height, (right + 1) / width, (bottom + 1) / height),
        "accent": accent,
        "luminance": luminance,
    }


def _clear_area(slide) -> tuple[tuple[int, int, int, int], tuple[int, int, int], bool]:
    """Region of the slide free of artwork, an accent colour, and whether the ground is dark."""
    prs = slide.part.package.presentation_part.presentation
    sw, sh = prs.slide_width, prs.slide_height
    left, top, right, bottom = 0.0, 0.0, float(sw), float(sh)
    accent = None
    dark = False
    for blob, x, y, w, h in _background_blobs(slide):
        result = _analyse_artwork(blob)
        if not result:
            continue
        l, t, r, b = result["box"]
        left, top = max(left, x + l * w), max(top, y + t * h)
        right, bottom = min(right, x + r * w), min(bottom, y + b * h)
        accent = accent or result["accent"]
        dark = dark or result["luminance"] < 110
        break
    # Vector artwork hugging an edge (waves, side bars, header/footer strips).
    for holder in (slide, slide.slide_layout, slide.slide_layout.slide_master):
        for shape in holder.shapes:
            if shape.is_placeholder or shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                continue
            if shape.has_text_frame and shape.text_frame.text.strip():
                continue
            if shape.width * shape.height >= 0.85 * sw * sh:
                continue
            s_right, s_bottom = shape.left + shape.width, shape.top + shape.height
            if shape.left <= 0.02 * sw and shape.width < 0.5 * sw and shape.height > 0.3 * sh:
                left = max(left, s_right)
            elif s_right >= 0.98 * sw and shape.width < 0.5 * sw and shape.height > 0.3 * sh:
                right = min(right, shape.left)
            elif shape.top <= 0.02 * sh and shape.height < 0.35 * sh and shape.width > 0.3 * sw:
                top = max(top, s_bottom)
            elif s_bottom >= 0.98 * sh and shape.height < 0.35 * sh and shape.width > 0.3 * sw:
                bottom = min(bottom, shape.top)
    left, top = left + 0.03 * sw, top + 0.03 * sh
    right, bottom = right - 0.04 * sw, bottom - 0.05 * sh
    left, top = max(left, 0.05 * sw), max(top, 0.05 * sh)
    right, bottom = min(right, 0.95 * sw), min(bottom, 0.94 * sh)
    if right - left < 0.4 * sw or bottom - top < 0.4 * sh:
        # Artwork covers too much to trust: use plain margins.
        left, top, right, bottom = 0.07 * sw, 0.06 * sh, 0.93 * sw, 0.94 * sh
    return (int(left), int(top), int(right), int(bottom)), accent or (31, 41, 68), dark


def _add_text_areas(slide) -> dict:
    """Give a design-only slide a title and a content area inside its clear region."""
    (left, top, right, bottom), accent, dark = _clear_area(slide)
    prs = slide.part.package.presentation_part.presentation
    for placeholder in list(slide.placeholders):
        if placeholder.has_text_frame and not placeholder.text_frame.text.strip():
            placeholder._element.getparent().remove(placeholder._element)
    title_height = int(0.19 * prs.slide_height)
    body_top = top + title_height + int(0.02 * prs.slide_height)
    ink = RGBColor(255, 255, 255) if dark else RGBColor(31, 41, 55)
    heading = RGBColor(255, 255, 255) if dark else RGBColor(*(int(c * 0.65) for c in accent))

    title = slide.shapes.add_textbox(left, top, right - left, title_height)
    title.name = "Layah Title"
    title.text_frame.word_wrap = True
    title.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    title.text_frame.vertical_anchor = MSO_ANCHOR.MIDDLE
    run = title.text_frame.paragraphs[0].add_run()
    run.text = "Slide title"
    run.font.size = Pt(28)
    run.font.bold = True
    run.font.color.rgb = heading

    body = slide.shapes.add_textbox(left, body_top, right - left, bottom - body_top)
    body.name = "Layah Content"
    body.text_frame.word_wrap = True
    body.text_frame.auto_size = MSO_AUTO_SIZE.NONE
    body.text_frame.vertical_anchor = MSO_ANCHOR.TOP
    paragraph = body.text_frame.paragraphs[0]
    paragraph.space_after = Pt(6)
    run = paragraph.add_run()
    run.text = "Lesson content"
    run.font.size = Pt(20)
    run.font.color.rgb = ink
    return {"safe": (left, top, right, bottom)}


def _prepare_design_slides(sources) -> dict[int, dict]:
    """Add text areas to source slides that are artwork only. Returns slide_id -> layout info."""
    prepared: dict[int, dict] = {}
    for slide in sources:
        try:
            _editable_slots(slide)
            continue
        except TemplateIncompatible:
            pass
        if _is_design_only(slide):
            prepared[slide.slide_id] = _add_text_areas(slide)
    return prepared


def _place_design_image(target, body_shape, image_bytes: bytes, safe) -> None:
    """Put a lesson picture beside the text on a design-only slide, keeping its proportions."""
    from PIL import Image
    prs = target.part.package.presentation_part.presentation
    left, top, right, bottom = safe
    width, height = Image.open(io.BytesIO(image_bytes)).size
    body_shape.width = int((right - left) * IMAGE_TEXT_FACTOR)
    x0 = body_shape.left + body_shape.width + int(0.02 * prs.slide_width)
    avail_w, avail_h = right - x0, body_shape.height
    if avail_w < 0.15 * prs.slide_width:
        raise TemplateIncompatible("REQUIRED_IMAGE_SLOT_MISSING", "The lesson includes an image, but this slide design has no clear place for it.")
    scale = min(avail_w / width, avail_h / height)
    w, h = int(width * scale), int(height * scale)
    target.shapes.add_picture(io.BytesIO(image_bytes), x0 + (avail_w - w) // 2, body_shape.top + max(0, (avail_h - h) // 6), w, h)


def _font_size(shape, default: float) -> float:
    for paragraph in shape.text_frame.paragraphs:
        for run in paragraph.runs:
            if run.font.size:
                return run.font.size.pt
        if paragraph.font.size:
            return paragraph.font.size.pt
    return default


def _capacity(shape, default_size: float, width_factor: float = 1.0, size: float | None = None) -> tuple[int, int, int] | None:
    """Conservative character budget, with wrapping/paragraph overhead.

    This is a guard, not a claim to reproduce PowerPoint's text layout.
    """
    size = size or _font_size(shape, default_size)
    width_pt = shape.width * width_factor / 12700
    height_pt = shape.height / 12700
    tf = shape.text_frame
    width_pt -= (tf.margin_left + tf.margin_right) / 12700
    height_pt -= (tf.margin_top + tf.margin_bottom) / 12700
    if size < 12 or width_pt < 80 or height_pt < size * 1.25:
        return None
    chars_per_line = max(1, int(width_pt / (size * 0.68)))
    lines = max(0, int(height_pt / (size * 1.45)))
    return chars_per_line * lines, chars_per_line, lines


def _split_body(text: str, chars_per_line: int, lines_per_slide: int) -> list[str]:
    if not text.strip():
        return [""]
    paragraphs = [p.strip() for p in text.replace("\r", "").split("\n") if p.strip()]
    chunks: list[list[str]] = [[]]
    used = 0
    for paragraph in paragraphs:
        needed = max(1, math.ceil(len(paragraph) / chars_per_line)) + 1
        if needed > lines_per_slide:
            # Split long paragraphs at word boundaries, retaining every word.
            words = paragraph.split()
            pieces: list[str] = []
            current = ""
            limit = chars_per_line * max(1, lines_per_slide - 1)
            for word in words:
                if len(word) > chars_per_line:
                    raise TemplateIncompatible("TEXT_OVERFLOW", "A word is too long for the template's content area.")
                if current and len(current) + 1 + len(word) > limit:
                    pieces.append(current)
                    current = word
                else:
                    current = f"{current} {word}".strip()
            if current:
                pieces.append(current)
        else:
            pieces = [paragraph]
        for piece in pieces:
            piece_lines = max(1, math.ceil(len(piece) / chars_per_line)) + 1
            if used + piece_lines > lines_per_slide and chunks[-1]:
                chunks.append([])
                used = 0
            chunks[-1].append(piece)
            used += piece_lines
    return ["\n".join(chunk) for chunk in chunks]


def _set_text_preserving_style(shape, text: str) -> None:
    frame = shape.text_frame
    samples = [copy.deepcopy(paragraph._p) for paragraph in frame.paragraphs]
    for paragraph in list(frame.paragraphs):
        paragraph._p.getparent().remove(paragraph._p)
    lines = text.split("\n") or [""]
    for index, line in enumerate(lines):
        sample = samples[min(index, len(samples) - 1)]
        p = copy.deepcopy(sample)
        for child in list(p):
            if child.tag != qn("a:pPr") and child.tag != qn("a:endParaRPr"):
                p.remove(child)
        r = etree.Element(qn("a:r"))
        source_run = sample.find(qn("a:r"))
        if source_run is not None and source_run.find(qn("a:rPr")) is not None:
            r.append(copy.deepcopy(source_run.find(qn("a:rPr"))))
        t = etree.SubElement(r, qn("a:t"))
        t.text = line
        p.insert(1 if len(p) and p[0].tag == qn("a:pPr") else 0, r)
        frame._txBody.append(p)
    frame.word_wrap = True
    frame.auto_size = MSO_AUTO_SIZE.NONE


def _clone_slide(prs, source):
    target = prs.slides.add_slide(source.slide_layout)
    target._element.attrib.clear()
    target._element.attrib.update(source._element.attrib)
    for element in list(target._element):
        target._element.remove(element)
    for element in source._element:
        target._element.append(copy.deepcopy(element))
    # python-pptx caches its shape collection when add_slide creates default
    # placeholders. Point it at the copied slide tree instead.
    target.__dict__.pop("shapes", None)
    id_map = {}
    for rel in source.part.rels.values():
        if rel.reltype.endswith("/slideLayout") or rel.reltype.endswith("/notesSlide"):
            continue
        new_id = target.part.relate_to(rel.target_ref if rel.is_external else rel.target_part, rel.reltype, is_external=rel.is_external)
        id_map[rel.rId] = new_id
    for element in target._element.iter():
        for name, value in list(element.attrib.items()):
            if name.startswith("{" + REL_NS + "}") and value in id_map:
                element.set(name, id_map[value])
    return target


def _remove_slide(prs, index: int):
    slide_id = prs.slides._sldIdLst[index]
    prs.part.drop_rel(slide_id.rId)
    prs.slides._sldIdLst.remove(slide_id)


def _replace_picture(slide, image_bytes: bytes | None):
    slide_area = slide.part.package.presentation_part.presentation.slide_width * slide.part.package.presentation_part.presentation.slide_height
    pictures = [s for s in slide.shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE and 0.12 < s.width * s.height / slide_area < 0.75]
    if not pictures:
        if image_bytes:
            placeholders = [
                s for s in slide.placeholders
                if "PICTURE" in str(s.placeholder_format.type).upper()
                and s.width * s.height > slide.part.package.presentation_part.presentation.slide_width * slide.part.package.presentation_part.presentation.slide_height * 0.12
            ]
            if len(placeholders) == 1:
                placeholders[0].insert_picture(io.BytesIO(image_bytes))
                return True
            raise TemplateIncompatible("REQUIRED_IMAGE_SLOT_MISSING", "The lesson includes an image, but this slide design has no clear place for it.")
        return False
    if len(pictures) > 1:
        raise TemplateIncompatible("AMBIGUOUS_MAPPING", "This slide has multiple large pictures, so the lesson image destination is unclear.")
    picture = pictures[0]
    if image_bytes:
        part, rid = slide.part.get_or_add_image_part(io.BytesIO(image_bytes))
        picture._element.blipFill.blip.set(qn("r:embed"), rid)
    else:
        picture._element.getparent().remove(picture._element)
    return True


def _set_notes(slide, text: str):
    if text:
        slide.notes_slide.notes_text_frame.text = text


def _choose_source(sources, index: int, total: int, title: str, body: str, design: dict | None = None, has_image: bool = False):
    design = design or {}
    if len(sources) >= total:
        alternatives = sources[1:-1] if len(sources) > 2 else sources
        candidates = [sources[index]] + [source for source in alternatives if source is not sources[index]]
    elif index == 0:
        candidates = [sources[0]]
    elif index == total - 1 and len(sources) > 2:
        candidates = [sources[-1]]
    else:
        content_sources = sources[1:-1] if len(sources) > 2 else sources
        offset = (index - 1) % len(content_sources)
        candidates = content_sources[offset:] + content_sources[:offset]
    last_error = None
    for candidate in candidates:
        try:
            lenient = candidate.slide_id in design
            title_shape, body_shape = _editable_slots(candidate, lenient=lenient)
            title_capacity = _capacity(title_shape, 30)
            if not title_capacity or len(title) > title_capacity[0]:
                raise TemplateIncompatible("TEXT_OVERFLOW", "The title or content area is too small.")
            factor = IMAGE_TEXT_FACTOR if lenient and has_image else 1.0
            body_size = None
            if lenient:
                # Use the largest size that keeps the section on one or two slides.
                parts = None
                for size in DESIGN_BODY_SIZES:
                    body_capacity = _capacity(body_shape, 18, factor, size)
                    if not body_capacity:
                        continue
                    try:
                        parts = _split_body(body, body_capacity[1], body_capacity[2])
                    except TemplateIncompatible as exc:
                        last_error = exc
                        continue
                    body_size = size
                    if len(parts) <= 2:
                        break
                if parts is None:
                    raise last_error or TemplateIncompatible("TEXT_OVERFLOW", "The title or content area is too small.")
                max_parts = MAX_PARTS_DESIGN_ONLY
            else:
                body_capacity = _capacity(body_shape, 18)
                if not body_capacity:
                    raise TemplateIncompatible("TEXT_OVERFLOW", "The title or content area is too small.")
                parts = _split_body(body, body_capacity[1], body_capacity[2])
                max_parts = MAX_PARTS_PER_SECTION
            if len(parts) > max_parts:
                raise TemplateIncompatible("TEXT_OVERFLOW", f"The lesson section needs more than {max_parts} slides in this design.")
            if len(parts) > 1 and len(f"{title} (continued)") > title_capacity[0]:
                raise TemplateIncompatible("TEXT_OVERFLOW", "The continuation title is too long for this design.")
            return candidate, title_shape, body_shape, parts, body_size
        except TemplateIncompatible as exc:
            last_error = exc
    raise last_error or TemplateIncompatible("NO_EDITABLE_CONTENT", "No suitable slide design was found.")


def render_template(data: bytes, slides: list[dict[str, Any]], images: dict[int, bytes] | None = None) -> bytes:
    validate_source(data)
    if not slides or len(slides) > 30:
        raise TemplateIncompatible("INVALID_CONTENT", "The lesson has an invalid number of slides.")
    try:
        prs = Presentation(io.BytesIO(data))
    except Exception as exc:
        raise TemplateIncompatible("INVALID_TEMPLATE", "This PowerPoint could not be opened.") from exc
    sources = list(prs.slides)
    if not sources:
        raise TemplateIncompatible("NO_EDITABLE_CONTENT", "The PowerPoint has no example slides. Add editable title and content slides, then upload it again.")
    if len(sources) > 30:
        raise TemplateIncompatible("INVALID_TEMPLATE", "Use a PowerPoint with at most 30 example slides.")
    images = images or {}
    design = _prepare_design_slides(sources)
    generated = 0
    expected_bodies: list[tuple[int, str]] = []
    for index, item in enumerate(slides):
        try:
            title = str(item.get("title", "")).strip()
            body = str(item.get("content", "")).strip()
            source, src_title, src_body, parts, body_size = _choose_source(
                sources, index, len(slides), title, body, design, bool(images.get(index)),
            )
            if re.sub(r"\s+", " ", " ".join(parts)).strip() != re.sub(r"\s+", " ", body).strip():
                raise TemplateIncompatible("VALIDATION_FAILED", "A lesson section could not be divided without losing content.")
            for part_index, part_text in enumerate(parts):
                if generated >= MAX_PHYSICAL_SLIDES:
                    raise TemplateIncompatible("TEXT_OVERFLOW", "This template would need too many slides.")
                target = _clone_slide(prs, source)
                target_title = next(s for s in target.shapes if s.shape_id == src_title.shape_id)
                target_body = next(s for s in target.shapes if s.shape_id == src_body.shape_id)
                part_title = title if part_index == 0 else f"{title} (continued)"
                _set_text_preserving_style(target_title, part_title)
                _set_text_preserving_style(target_body, part_text)
                expected_bodies.append((target_body.shape_id, part_text))
                lesson_image = images.get(index) if part_index == 0 else None
                if body_size is not None:
                    # Design-only template: text areas were added by Layah.
                    for paragraph in target_body.text_frame.paragraphs:
                        for run in paragraph.runs:
                            run.font.size = Pt(body_size)
                    if lesson_image:
                        _place_design_image(target, target_body, lesson_image, design[source.slide_id]["safe"])
                else:
                    _replace_picture(target, lesson_image)
                speaker_notes = str(item.get("speakerNotes", "")).strip()
                if part_index > 0 and speaker_notes:
                    speaker_notes = f"{speaker_notes}\n\nContinuation of this section."
                _set_notes(target, speaker_notes)
                generated += 1
        except TemplateIncompatible as exc:
            exc.slide = index + 1
            raise
    for index in range(len(sources) - 1, -1, -1):
        _remove_slide(prs, index)
    result = io.BytesIO()
    prs.save(result)
    output = result.getvalue()
    # The package must be reopenable and all logical text must survive export.
    check = Presentation(io.BytesIO(output))
    if len(check.slides) != generated:
        raise TemplateIncompatible("VALIDATION_FAILED", "The exported PowerPoint failed validation.")
    for slide, (shape_id, expected_text) in zip(check.slides, expected_bodies):
        actual = next((shape.text for shape in slide.shapes if shape.shape_id == shape_id), None)
        if actual != expected_text:
            raise TemplateIncompatible("VALIDATION_FAILED", "Some lesson content is missing from the exported PowerPoint.")
    return output
