import io
import json
import os
import struct
import unittest
import zlib
from zipfile import ZipFile
from unittest.mock import patch

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches

from template_engine import TemplateIncompatible, render_template


def source_deck():
    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[1])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = RGBColor(244, 245, 230)
    slide.shapes.title.text = "Old title"
    slide.placeholders[1].text = "Old lesson content"
    brand = slide.shapes.add_shape(1, Inches(0.1), Inches(0.1), Inches(0.2), Inches(0.2))
    brand.fill.solid()
    brand.fill.fore_color.rgb = RGBColor(20, 90, 150)
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def png(red: int, green: int, blue: int):
    def chunk(kind, data):
        return struct.pack(">I", len(data)) + kind + data + struct.pack(">I", zlib.crc32(kind + data))
    return (
        b"\x89PNG\r\n\x1a\n"
        + chunk(b"IHDR", struct.pack(">IIBBBBB", 2, 2, 8, 2, 0, 0, 0))
        + chunk(b"IDAT", zlib.compress((b"\x00" + bytes((red, green, blue)) * 2) * 2))
        + chunk(b"IEND", b"")
    )


def artwork_only_deck(slide_count: int = 1):
    """A school-style design: full-slide artwork with a wave down the left edge and a
    faint crest in the middle, plus a plain header band, and no text areas at all."""
    from PIL import Image, ImageDraw

    art = Image.new("RGBA", (800, 450), (255, 255, 255, 255))
    draw = ImageDraw.Draw(art)
    draw.polygon([(0, 0), (90, 0), (140, 450), (0, 450)], fill=(90, 145, 205, 255))
    draw.ellipse([350, 170, 450, 270], fill=(235, 238, 245, 255))  # faint crest, floating
    art_bytes = io.BytesIO()
    art.save(art_bytes, "PNG")
    band = io.BytesIO()
    Image.new("RGB", (200, 40), (255, 255, 255)).save(band, "PNG")

    deck = Presentation()
    deck.slide_width, deck.slide_height = Inches(13.333), Inches(7.5)
    for _ in range(slide_count):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        slide.shapes.add_picture(io.BytesIO(art_bytes.getvalue()), 0, 0, deck.slide_width, deck.slide_height)
        slide.shapes.add_picture(io.BytesIO(band.getvalue()), Inches(1.7), Inches(0.3), Inches(10.9), Inches(1.1))
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


def sectioned_deck():
    """A template the way PowerPoint saves one that uses sections: the presentation part
    carries a p14:sectionLst naming every slide."""
    from lxml import etree

    p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
    pml = "http://schemas.openxmlformats.org/presentationml/2006/main"
    deck = Presentation()
    for n in range(3):
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = f"Old title {n}"
        slide.placeholders[1].text = f"Old body {n}"
    root = deck.part._element
    ids = [e.get("id") for e in root.iter("{%s}sldId" % pml)]
    root.append(etree.fromstring(
        f'<p:extLst xmlns:p="{pml}"><p:ext uri="{{521415D9-36F7-43E2-AB2F-B90AF26B5E84}}">'
        f'<p14:sectionLst xmlns:p14="{p14}"><p14:section name="Default Section" id="{{11111111-2222-3333-4444-555555555555}}">'
        "<p14:sldIdLst>" + "".join(f'<p14:sldId id="{i}"/>' for i in ids) + "</p14:sldIdLst>"
        "</p14:section></p14:sectionLst></p:ext></p:extLst>"
    ))
    output = io.BytesIO()
    deck.save(output)
    return output.getvalue()


class UploadedTemplateTests(unittest.TestCase):
    def test_section_list_follows_the_new_slides(self):
        # Regression: the section list kept naming the deleted template slides and left every
        # generated slide outside any section, which makes PowerPoint offer to "repair" the file.
        from lxml import etree

        p14 = "http://schemas.microsoft.com/office/powerpoint/2010/main"
        pml = "http://schemas.openxmlformats.org/presentationml/2006/main"
        slides = [{"title": f"Lesson {n}", "content": f"Content {n}", "speakerNotes": ""} for n in range(5)]
        rendered = render_template(sectioned_deck(), slides)
        with ZipFile(io.BytesIO(rendered)) as archive:
            root = etree.fromstring(archive.read("ppt/presentation.xml"))
        live = [e.get("id") for e in root.iter("{%s}sldId" % pml)]
        in_sections = [e.get("id") for e in root.iter("{%s}sldId" % p14)]
        self.assertEqual(len(live), 5)
        self.assertEqual(sorted(in_sections), sorted(live))

    def test_control_characters_in_lesson_text_are_stripped(self):
        # Regression: text copied from PDFs carries form feeds/NULs; python-pptx raised
        # ValueError and anything that got through made the XML invalid.
        slides = [{"title": "Ti\x0ctle", "content": "Body \x00 with \x0b control \x1f chars", "speakerNotes": "no\x01tes"}]
        rendered = render_template(source_deck(), slides)
        result = Presentation(io.BytesIO(rendered))
        self.assertEqual(result.slides[0].shapes.title.text, "Title")
        self.assertEqual(result.slides[0].placeholders[1].text, "Body  with  control  chars")
        self.assertEqual(result.slides[0].notes_slide.notes_text_frame.text, "notes")

    def test_completed_source_reused_for_full_lesson_without_old_text(self):
        slides = [{"title": f"Lesson {n}", "content": f"New content {n}", "speakerNotes": "Teacher note"} for n in range(13)]
        source = source_deck()
        rendered = render_template(source, slides)
        with ZipFile(io.BytesIO(source)) as before, ZipFile(io.BytesIO(rendered)) as after:
            self.assertEqual(before.read("ppt/theme/theme1.xml"), after.read("ppt/theme/theme1.xml"))
        result = Presentation(io.BytesIO(rendered))
        self.assertEqual(len(result.slides), 13)
        self.assertEqual(result.slides[-1].shapes.title.text, "Lesson 12")
        self.assertEqual(result.slides[-1].placeholders[1].text, "New content 12")
        self.assertEqual(result.slides[0].notes_slide.notes_text_frame.text, "Teacher note")
        all_text = " ".join(shape.text for slide in result.slides for shape in slide.shapes if shape.has_text_frame)
        self.assertNotIn("Old lesson content", all_text)
        self.assertEqual(result.slides[0].shapes[-1].fill.fore_color.rgb, RGBColor(20, 90, 150))
        self.assertEqual(result.slides[0].background.fill.fore_color.rgb, RGBColor(244, 245, 230))

    def test_very_dense_content_returns_fit_warning(self):
        with self.assertRaises(TemplateIncompatible) as raised:
            render_template(source_deck(), [{"title": "Dense lesson", "content": "example " * 5000}])
        self.assertEqual(raised.exception.code, "TEXT_OVERFLOW")
        self.assertEqual(raised.exception.slide, 1)

    def test_extra_old_lesson_text_is_rejected(self):
        deck = Presentation(io.BytesIO(source_deck()))
        deck.slides[0].shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text = "Old lesson that could leak"
        output = io.BytesIO()
        deck.save(output)
        with self.assertRaises(TemplateIncompatible) as raised:
            render_template(output.getvalue(), [{"title": "New", "content": "Content"}])
        self.assertEqual(raised.exception.code, "AMBIGUOUS_MAPPING")

    def test_lesson_picture_replaces_old_picture_and_keeps_branding(self):
        deck = Presentation(io.BytesIO(source_deck()))
        deck.slides[0].shapes.add_picture(io.BytesIO(png(255, 0, 0)), Inches(7), Inches(1.5), width=Inches(4), height=Inches(3))
        source = io.BytesIO()
        deck.save(source)
        result = Presentation(io.BytesIO(render_template(source.getvalue(), [{"title": "New", "content": "Content"}], {0: png(0, 0, 255)})))
        self.assertEqual(result.slides[0].shapes[-1].image.blob, png(0, 0, 255))
        self.assertEqual(result.slides[0].shapes[-2].fill.fore_color.rgb, RGBColor(20, 90, 150))

    def test_generated_image_without_template_frame_returns_warning(self):
        with self.assertRaises(TemplateIncompatible) as raised:
            render_template(source_deck(), [{"title": "New", "content": "Content"}], {0: png(0, 0, 255)})
        self.assertEqual(raised.exception.code, "REQUIRED_IMAGE_SLOT_MISSING")

    def test_uses_another_source_slide_when_first_design_is_ambiguous(self):
        deck = Presentation(io.BytesIO(source_deck()))
        deck.slides[0].shapes.add_textbox(Inches(1), Inches(2), Inches(3), Inches(1)).text = "Old extra content"
        second = deck.slides.add_slide(deck.slide_layouts[1])
        second.shapes.title.text = "Other old title"
        second.placeholders[1].text = "Other old content"
        output = io.BytesIO()
        deck.save(output)
        result = Presentation(io.BytesIO(render_template(output.getvalue(), [{"title": "New title", "content": "New content"}])))
        self.assertEqual(result.slides[0].shapes.title.text, "New title")
        self.assertEqual(result.slides[0].placeholders[1].text, "New content")

    def test_artwork_only_single_slide_template_generates_every_slide(self):
        slides = [{"title": f"Lesson {n}", "content": f"Point one for {n}\nPoint two for {n}", "speakerNotes": "note"} for n in range(13)]
        result = Presentation(io.BytesIO(render_template(artwork_only_deck(), slides)))
        self.assertEqual(len(result.slides), 13)
        for n, slide in enumerate(result.slides):
            titles = [s for s in slide.shapes if s.name == "Layah Title"]
            bodies = [s for s in slide.shapes if s.name == "Layah Content"]
            self.assertEqual(titles[0].text_frame.text, f"Lesson {n}")
            self.assertEqual(bodies[0].text_frame.text, f"Point one for {n}\nPoint two for {n}")
        # Design pictures stay, and the text sits clear of the wave down the left edge.
        first = result.slides[0]
        self.assertEqual(sum(1 for s in first.shapes if s.shape_type == 13), 2)
        wave_edge = int(result.slide_width * 140 / 800)
        for shape in first.shapes:
            if shape.name.startswith("Layah"):
                self.assertGreater(shape.left, wave_edge)
                self.assertLessEqual(shape.left + shape.width, result.slide_width)
        self.assertEqual(first.notes_slide.notes_text_frame.text, "note")

    def test_artwork_only_template_places_lesson_image_beside_text(self):
        result = Presentation(io.BytesIO(render_template(
            artwork_only_deck(), [{"title": "With picture", "content": "Some content"}], {0: png(0, 0, 255)},
        )))
        slide = result.slides[0]
        body = next(s for s in slide.shapes if s.name == "Layah Content")
        picture = slide.shapes[-1]
        self.assertEqual(picture.image.blob, png(0, 0, 255))
        self.assertGreaterEqual(picture.left, body.left + body.width)
        self.assertLessEqual(picture.left + picture.width, result.slide_width)

    def test_artwork_only_template_splits_long_content_across_slides(self):
        content = "\n".join(f"Line {n}: keep both sides of the equation balanced at every step" for n in range(40))
        result = Presentation(io.BytesIO(render_template(artwork_only_deck(), [{"title": "Long", "content": content}])))
        self.assertGreater(len(result.slides), 1)
        joined = " ".join(s.text_frame.text for slide in result.slides for s in slide.shapes if s.name == "Layah Content")
        self.assertIn("Line 39", joined)

    def test_artwork_only_template_with_text_is_still_checked_strictly(self):
        deck = Presentation(io.BytesIO(artwork_only_deck()))
        deck.slides[0].shapes.add_textbox(Inches(2), Inches(2), Inches(4), Inches(1)).text = "Old lesson text"
        output = io.BytesIO()
        deck.save(output)
        with self.assertRaises(TemplateIncompatible):
            render_template(output.getvalue(), [{"title": "New", "content": "Content"}])

    def test_service_requires_secret_and_returns_structured_fit_issue(self):
        from main import app

        client = app.test_client()
        with patch.dict(os.environ, {"PPT_TEMPLATE_SERVICE_SECRET": "test-secret"}):
            unauthenticated = client.post("/render-uploaded-template")
            self.assertEqual(unauthenticated.status_code, 401)
            response = client.post(
                "/render-uploaded-template",
                headers={"X-Template-Service-Secret": "test-secret"},
                data={
                    "template": (io.BytesIO(source_deck()), "example.pptx"),
                    "slides": json.dumps([{"title": "Dense", "content": "many words " * 5000}]),
                },
            )
            self.assertEqual(response.status_code, 422)
            self.assertEqual(response.json["code"], "TEXT_OVERFLOW")
            self.assertEqual(response.json["slide"], 1)


if __name__ == "__main__":
    unittest.main()
